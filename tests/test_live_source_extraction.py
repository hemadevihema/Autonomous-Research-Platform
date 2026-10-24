"""Optional end-to-end check against a real open-access PMC article.

Run with RUN_LIVE_SOURCE_TESTS=1 to exercise live retrieval and both exporters.
"""
import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import openpyxl
from docx import Document

from agents.extraction import EvidenceExtractionAgent
from agents.qc import QualityControlAgent
from agents.scouting import ResearchScoutingAgent
from agents.validation import SourceValidationAgent
from exporters.docx_exporter import WordDocumentExporter
from exporters.excel_exporter import ExcelSpreadsheetExporter
from storage.database import ResearchDatabase


@unittest.skipUnless(os.environ.get("RUN_LIVE_SOURCE_TESTS") == "1", "set RUN_LIVE_SOURCE_TESTS=1 for live PMC retrieval")
class TestLiveSourceExtraction(unittest.TestCase):
    def test_real_pmc_source_through_retrieval_extraction_and_exports(self):
        test_root = Path(__file__).resolve().parents[1] / ".test_tmp"
        test_root.mkdir(exist_ok=True)
        with tempfile.TemporaryDirectory(dir=test_root) as tmp:
            root = Path(tmp)
            db = ResearchDatabase(root / "live.sqlite")
            scout = ResearchScoutingAgent(db)
            chunks = scout._fetch_pmc_full_text("PMC11097587")
            self.assertTrue(chunks, "NCBI returned no full-text sections for the known open-access PMC article")
            metadata = scout._get_cached_pmc_metadata("PMC11097587")
            self.assertTrue(metadata.get("source_title"), "Title should come from the original PMC XML metadata")
            self.assertIsInstance(metadata.get("publication_year"), int, "Publication year should come from source metadata")
            full_text = "\n".join(chunk["text"] for chunk in chunks)
            period = scout._detect_data_period_details(chunks, "full_text")
            source = {
                "id": "SRC-PMC11097587",
                "source_title": metadata["source_title"],
                "name_of_source": metadata.get("name_of_source") or "PMC article",
                "source_type": "Research article",
                "publication_year": metadata["publication_year"],
                "data_period": period["data_period"],
                "data_period_start_year": period["start_year"],
                "data_period_end_year": period["end_year"],
                "data_period_type": period["period_type"],
                "data_period_evidence": period["evidence"],
                "data_period_location": period["location"],
                "data_period_status": period["status"],
                "content_completeness": "full_text",
                "source_link": "https://pmc.ncbi.nlm.nih.gov/articles/PMC11097587/",
                "country": "India",
                "disease": "Obesity",
                "sections_covered": ["Epidemiology"],
                "content_completeness": "full_text",
                "content_chunks": chunks,
                "full_text_or_abstract": full_text,
            }
            result = EvidenceExtractionAgent(db).extract_evidence({
                "project_id": "LIVE-PMC-TEST",
                "validated_sources": [source],
                "audit_trail": [],
            })
            self.assertGreater(len(result["insights"]), 1)
            self.assertGreater(len(result["verbatims"]), 1)
            for item in result["verbatims"]:
                self.assertTrue(item["match_verified"])
                self.assertIn(item["exact_quote"], full_text)

            validation = SourceValidationAgent(db).validate_sources({
                "project_id": "LIVE-PMC-TEST",
                "candidate_sources": [source],
                "min_publication_year": 2021,
                "data_period_start_year": 2021,
                "data_period_end_year": 2026,
                "countries": ["India"],
                "disease": "Obesity",
            })
            qc_state = {
                "countries": ["India"], "sections": ["Epidemiology"],
                "validated_sources": validation["validated_sources"],
                "review_sources": validation["review_sources"],
                "rejected_sources": validation["rejected_sources"],
                "data_period_start_year": 2021, "data_period_end_year": 2026,
                "min_publication_year": 2021, "recovery_attempts": 0,
                "insights": [], "verbatims": [],
            }
            qc_result = QualityControlAgent().evaluate_project(qc_state)
            covered_years = sum(
                cell["status"] == "Covered"
                for cell in qc_result["year_wise_coverage"].values()
            )
            self.assertLessEqual(covered_years, 6)

            with patch("exporters.excel_exporter.OUTPUT_DIR", root), patch("exporters.docx_exporter.OUTPUT_DIR", root):
                xlsx_path = ExcelSpreadsheetExporter().export(
                    "LIVE-PMC-TEST", "Obesity in India", validation["validated_sources"],
                    qc_report=qc_result["qc_report"],
                    insights=[], verbatims=[], data_period_sources=validation["candidate_sources"],
                )
                docx_path = WordDocumentExporter().export(
                    project_id="LIVE-PMC-TEST", topic="Obesity in India", disease="Obesity",
                    countries=["India"], sections=["Epidemiology"],
                    validated_sources=validation["validated_sources"], insights=[],
                    verbatims=[], qc_report=qc_result["qc_report"],
                    extraction_completeness=result["extraction_completeness"],
                    min_publication_year=2021, data_period_start_year=2021,
                    data_period_end_year=2026, review_sources=validation["review_sources"],
                )
            workbook = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
            master = workbook["Research Sources"]
            self.assertEqual(master.max_column, 14)
            self.assertEqual(master.cell(1, 6).value, "Period of the data used")
            period_audit = workbook["Data Period Audit"]
            self.assertEqual(period_audit.cell(2, 5).value, period["data_period"])
            audited_source = (validation["validated_sources"] + validation["review_sources"])[0]
            self.assertEqual(period_audit.cell(2, 9).value, audited_source["data_period_status"])
            qc_sheet = workbook["QC Audit & Coverage"]
            self.assertTrue(any(cell.value == "India | Epidemiology | 2021" for row in qc_sheet.iter_rows() for cell in row))
            workbook.close()

            document = Document(docx_path)
            exported_text = "\n".join(p.text for p in document.paragraphs)
            exported_text += "\n" + "\n".join(
                cell.text for table in document.tables for row in table.rows for cell in row.cells
            )
            self.assertIn(period["data_period"], exported_text)
            self.assertIn("Data-period coverage by year", exported_text)
            self.assertEqual(len(qc_result["year_wise_coverage"]), 6)


if __name__ == "__main__":
    unittest.main()
