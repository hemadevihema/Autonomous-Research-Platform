"""Focused regression tests for multi-finding, source-grounded extraction."""
import tempfile
import unittest
from pathlib import Path

from docx import Document
import openpyxl

from agents.extraction import EvidenceExtractionAgent
from exporters.docx_exporter import WordDocumentExporter
from exporters.excel_exporter import ExcelSpreadsheetExporter
from storage.database import ResearchDatabase


FIVE_FINDINGS = [
    "Among 5,120 adults sampled across 31 Indian states, obesity prevalence was 24.8%.",
    "Women had a higher prevalence of abdominal obesity than men across every age group.",
    "Urban residence was associated with greater odds of obesity than rural residence (adjusted odds ratio 1.42).",
    "Prevalence increased with age and peaked among adults aged 50–59 years.",
    "The study estimated regional prevalence at 19.2% in rural communities and 27.5% in urban communities.",
]


def make_source(chunks=None, **overrides):
    source = {
        "id": "SRC-TEST-MULTI",
        "source_title": "National obesity prevalence study in India",
        "name_of_source": "Synthetic test fixture (not real research)",
        "source_type": "Research article",
        "publication_year": 2024,
        "data_period": "2021–2022",
        "source_link": "https://example.invalid/test-source",
        "country": "India",
        "disease": "Obesity",
        "sections_covered": ["Epidemiology"],
        "validation_status": "Passed initial checks",
        "content_completeness": "full_text",
        "content_chunks": chunks or [{"location": "Results", "text": " ".join(FIVE_FINDINGS)}],
        "full_text_or_abstract": " ".join(FIVE_FINDINGS),
    }
    source.update(overrides)
    return source


class TestMultiEvidenceExtraction(unittest.TestCase):
    def setUp(self):
        # Keep test artifacts under the writable project workspace; the system
        # temp directory may be protected by the managed Windows sandbox.
        test_root = Path(__file__).resolve().parents[1] / ".test_tmp"
        test_root.mkdir(exist_ok=True)
        self.temp_dir = tempfile.TemporaryDirectory(dir=test_root)
        self.db = ResearchDatabase(Path(self.temp_dir.name) / "test.sqlite")
        self.extractor = EvidenceExtractionAgent(self.db)

    def tearDown(self):
        self.temp_dir.cleanup()

    def run_extraction(self, source):
        return self.extractor.extract_evidence({
            "project_id": "TEST-MULTI",
            "validated_sources": [source],
            "audit_trail": [],
        })

    def test_five_distinct_findings_and_quotes_are_all_retained(self):
        result = self.run_extraction(make_source())
        self.assertEqual(len(result["insights"]), 5)
        self.assertEqual(len(result["verbatims"]), 5)
        self.assertEqual(len({item["insight_text"] for item in result["insights"]}), 5)
        self.assertTrue(all(item["match_verified"] for item in result["verbatims"]))
        self.assertEqual(result["extraction_completeness"]["SRC-TEST-MULTI"]["status"], "complete")

    def test_exact_quote_checker_rejects_invented_text_and_accepts_whitespace_changes(self):
        text = "The obesity prevalence was 24.8% among adults."
        self.assertTrue(self.extractor.verify_exact_quote("The obesity\nprevalence was 24.8% among adults.", text))
        self.assertFalse(self.extractor.verify_exact_quote("The prevalence was 99% among adults.", text))

    def test_generic_fallback_text_is_not_extracted_or_verified(self):
        fake = "National Library of Medicine indexed research on Obesity in India regarding Epidemiology."
        result = self.run_extraction(make_source(
            chunks=[{"location": "Abstract", "text": fake}],
            full_text_or_abstract=fake,
            content_completeness="abstract_only",
        ))
        self.assertEqual(result["insights"], [])
        self.assertEqual(result["verbatims"], [])

    def test_duplicate_passages_across_chunks_are_deduplicated(self):
        repeated = FIVE_FINDINGS[0]
        result = self.run_extraction(make_source(chunks=[
            {"location": "Abstract", "text": repeated},
            {"location": "Results", "text": repeated},
        ], full_text_or_abstract=f"{repeated}\n{repeated}"))
        self.assertEqual(len(result["insights"]), 1)
        self.assertEqual(len(result["verbatims"]), 1)

    def test_findings_from_multiple_chunks_are_aggregated(self):
        chunks = [
            {"location": "Results, page 2", "text": " ".join(FIVE_FINDINGS[:2])},
            {"location": "Results, page 3", "text": " ".join(FIVE_FINDINGS[2:])},
        ]
        result = self.run_extraction(make_source(chunks=chunks))
        self.assertEqual(len(result["insights"]), 5)
        self.assertEqual(len(result["verbatims"]), 5)
        self.assertTrue(any(item["location_reference"] == "Results, page 3" for item in result["verbatims"]))

    def test_partial_abstract_reports_shortfall_without_padding(self):
        passage = "Obesity prevalence was 24.8% among adults in the national sample."
        source = make_source(
            chunks=[{"location": "Abstract", "text": passage}],
            full_text_or_abstract=passage,
            content_completeness="abstract_only",
        )
        result = self.run_extraction(source)
        self.assertEqual(len(result["insights"]), 1)
        self.assertEqual(len(result["verbatims"]), 1)
        self.assertEqual(result["extraction_completeness"][source["id"]]["status"], "below_target")
        self.assertIn("Abstract-only", result["insights"][0]["limitations"])

    def test_records_persist_individually(self):
        result = self.run_extraction(make_source())
        self.assertEqual(len(self.db.get_insights_for_project("TEST-MULTI")), 5)
        self.assertEqual(len(self.db.get_verbatims_for_project("TEST-MULTI")), 5)
        self.assertEqual(len(result["validated_sources"][0]["insights"].split("Insight ")) - 1, 5)

    def test_excel_serializes_every_finding_and_preserves_master_headers(self):
        result = self.run_extraction(make_source())
        source = result["validated_sources"][0]
        source = dict(source)
        source["verbatim"] = "\n\n".join(
            f'Quote {idx}: "{item["exact_quote"]}"'
            for idx, item in enumerate(result["verbatims"][:3], start=1)
        )
        path = ExcelSpreadsheetExporter().export(
            "TEST-MULTI-EVIDENCE",
            "Obesity in India",
            [source],
            qc_report={"overall_status": "PASSED_WITH_LIMITATIONS", "issues": []},
            insights=result["insights"],
            verbatims=result["verbatims"][:3],
        )
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        master = workbook["Research Sources"]
        self.assertEqual(master.max_column, 14)
        self.assertEqual(master.cell(1, 12).value, "Insights")
        self.assertEqual(master.cell(1, 13).value, "Verbatim")
        self.assertEqual(master.cell(2, 12).value.count("Insight "), 5)
        self.assertEqual(master.cell(2, 13).value.count("Quote "), 3)
        evidence = workbook["Evidence Details"]
        self.assertEqual(evidence.max_row, 9)  # Five insight records plus three distinct quote records.
        record_types = [evidence.cell(row, 4).value for row in range(2, evidence.max_row + 1)]
        self.assertEqual(record_types.count("Insight"), 5)
        self.assertEqual(record_types.count("Verbatim"), 3)

    def test_word_report_contains_all_extracted_items(self):
        result = self.run_extraction(make_source())
        path = WordDocumentExporter().export(
            project_id="TEST-MULTI-WORD",
            topic="Obesity in India",
            disease="Obesity",
            countries=["India"],
            sections=["Epidemiology"],
            validated_sources=result["validated_sources"],
            insights=result["insights"],
            verbatims=result["verbatims"],
            qc_report={"overall_status": "PASSED", "missing_coverage": []},
            extraction_completeness=result["extraction_completeness"],
        )
        document = Document(path)
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        table_text = "\n".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
        combined = text + "\n" + table_text
        for finding in FIVE_FINDINGS:
            self.assertIn(finding, combined)


if __name__ == "__main__":
    unittest.main()
