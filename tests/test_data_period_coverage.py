"""Regression tests for verified collection periods and annual coverage."""
import tempfile
import unittest
from pathlib import Path

import openpyxl
from docx import Document

from agents.correction import AutonomousCorrectionEngine
from agents.qc import QualityControlAgent
from agents.scouting import ResearchScoutingAgent
from agents.validation import SourceValidationAgent
from config.settings import SPREADSHEET_COLUMNS
from exporters.docx_exporter import WordDocumentExporter
from exporters.excel_exporter import ExcelSpreadsheetExporter
from unittest.mock import patch


class TestDataPeriodCoverage(unittest.TestCase):
    def source(self, **overrides):
        record = {
            "id": "SRC-PERIOD-TEST",
            "source_title": "Obesity prevalence study in India",
            "name_of_source": "Synthetic test journal (not real research)",
            "publication_year": 2025,
            "data_period": "2021–2023",
            "data_period_start_year": 2021,
            "data_period_end_year": 2023,
            "data_period_type": "data_collection",
            "data_period_evidence": "Data were collected between 2021 and 2023.",
            "data_period_location": "Methods",
            "data_period_status": "Verified",
            "content_completeness": "full_text",
            "country": "India",
            "disease": "Obesity",
            "sections_covered": ["Epidemiology"],
            "full_text_or_abstract": (
                "Obesity prevalence was studied in India. Data were collected between 2021 and 2023."
            ),
            "source_link": "https://example.invalid/synthetic-test",
            "source_provider": "Test fixture",
            "provider_record_id": "fixture-period",
            "provider_record_url": "https://example.invalid/record/fixture-period",
            "source_authenticity_verified": True,
            "publication_year_verified": True,
            "validation_status": "Retrieved",
        }
        record.update(overrides)
        return record

    def test_2024_publication_keeps_2018_2020_collection_period(self):
        scout = ResearchScoutingAgent()
        details = scout._detect_data_period_details(
            [{"location": "Methods", "text": "Data were collected between 2018 and 2020. Published in 2024."}],
            "full_text",
        )
        self.assertEqual(details["data_period"], "2018–2020")
        self.assertEqual((details["start_year"], details["end_year"]), (2018, 2020))
        self.assertEqual(details["status"], "Verified")
        source = self.source(
            publication_year=2024, data_period=details["data_period"],
            data_period_start_year=details["start_year"], data_period_end_year=details["end_year"],
            data_period_evidence=details["evidence"], data_period_location=details["location"],
            data_period_status=details["status"], data_period_type=details["period_type"],
            full_text_or_abstract="Obesity prevalence in India. Data were collected between 2018 and 2020.",
        )
        checked = SourceValidationAgent().validate_sources({
            "candidate_sources": [source], "min_publication_year": 2021,
            "data_period_start_year": 2021, "data_period_end_year": 2026,
            "countries": ["India"], "disease": "Obesity",
        })
        self.assertEqual(len(checked["review_sources"]), 1)
        self.assertEqual(checked["review_sources"][0]["publication_year"], 2024)
        self.assertEqual(checked["review_sources"][0]["data_period"], "2018–2020")
        self.assertEqual(checked["review_sources"][0]["data_period_status"], "Outside requested period")

    def test_2025_publication_records_2021_2023_period(self):
        scout = ResearchScoutingAgent()
        details = scout._detect_data_period_details(
            [{"location": "Methods", "text": "Survey data were collected from 2021 to 2023."}],
            "full_text",
        )
        self.assertEqual(details["data_period"], "2021–2023")
        self.assertEqual(details["period_type"], "survey")

    def test_no_period_evidence_remains_unknown_not_current_year(self):
        scout = ResearchScoutingAgent()
        details = scout._detect_data_period_details(
            [{"location": "Abstract", "text": "This article was published in 2026."}],
            "abstract_only",
        )
        self.assertEqual(details["data_period"], "Unknown")
        self.assertEqual(details["status"], "Unknown")
        self.assertIsNone(details["start_year"])
        source = self.source(
            publication_year=2026, data_period="2026", data_period_evidence="",
            data_period_start_year=None, data_period_end_year=None,
        )
        checked = SourceValidationAgent().validate_sources({
            "candidate_sources": [source], "min_publication_year": 2021,
            "data_period_start_year": 2021, "data_period_end_year": 2026,
            "countries": ["India"], "disease": "Obesity",
        })
        self.assertEqual(checked["review_sources"][0]["data_period"], "Unknown")
        self.assertEqual(checked["review_sources"][0]["data_period_status"], "Unknown")

    def test_qc_marks_only_years_2021_to_2023_covered(self):
        source = self.source()
        result = QualityControlAgent().evaluate_project({
            "countries": ["India"], "sections": ["Epidemiology"],
            "validated_sources": [source], "rejected_sources": [],
            "data_period_start_year": 2021, "data_period_end_year": 2026,
            "recovery_attempts": 0, "insights": [], "verbatims": [],
        })
        matrix = result["year_wise_coverage"]
        self.assertEqual([matrix[f"India | Epidemiology | {year}"]["status"] for year in range(2021, 2027)],
                         ["Covered", "Covered", "Covered", "No evidence", "No evidence", "No evidence"])
        self.assertEqual(len(result["year_coverage_gaps"]), 3)

    def test_missing_years_trigger_bounded_year_specific_searches(self):
        result = AutonomousCorrectionEngine().prepare_recovery({
            "disease": "Obesity", "min_publication_year": 2021,
            "missing_sections": [],
            "year_coverage_gaps": ["India | Epidemiology | 2024", "India | Epidemiology | 2025"],
            "recovery_attempts": 0, "audit_trail": [],
        })
        self.assertTrue(result["recovery_needed"])
        self.assertEqual([query["data_year"] for query in result["targeted_queries"]], [2024, 2025])
        self.assertIn("data collected 2024", result["targeted_queries"][0]["query"])

    def test_wrong_country_and_unsupported_section_are_not_assigned(self):
        source = self.source(
            source_title="Obesity prevalence study in Australia",
            full_text_or_abstract="Adult obesity prevalence in Australia was measured.",
            sections_covered=["Epidemiology", "Clinical Guidelines & Treatment Practices"],
        )
        result = SourceValidationAgent().validate_sources({
            "candidate_sources": [source], "min_publication_year": 2021,
            "data_period_start_year": 2021, "data_period_end_year": 2026,
            "disease": "Obesity", "countries": ["India"],
        })
        self.assertEqual(result["validated_sources"], [])
        self.assertEqual(result["rejected_sources"][0]["sections_covered"], ["Epidemiology"])
        self.assertTrue(any("assigned target country" in reason for reason in result["rejected_sources"][0]["validation_reasons"]))

    def test_xlsx_and_docx_report_periods_and_year_gaps(self):
        source = self.source()
        qc = QualityControlAgent().evaluate_project({
            "countries": ["India"], "sections": ["Epidemiology"],
            "validated_sources": [source], "rejected_sources": [],
            "data_period_start_year": 2021, "data_period_end_year": 2026,
            "recovery_attempts": 0, "insights": [], "verbatims": [],
        })["qc_report"]
        test_root = Path(__file__).resolve().parents[1] / ".test_tmp"
        test_root.mkdir(exist_ok=True)
        with tempfile.TemporaryDirectory(dir=test_root) as temp:
            with patch("exporters.excel_exporter.OUTPUT_DIR", Path(temp)), patch("exporters.docx_exporter.OUTPUT_DIR", Path(temp)):
                xlsx = ExcelSpreadsheetExporter().export("PERIOD-TEST", "Obesity", [source], qc_report=qc)
                docx = WordDocumentExporter().export(
                    project_id="PERIOD-TEST", topic="Obesity", disease="Obesity",
                    countries=["India"], sections=["Epidemiology"], validated_sources=[source],
                    insights=[], verbatims=[], qc_report=qc,
                    min_publication_year=2021, data_period_start_year=2021, data_period_end_year=2026,
                )
            workbook = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
            sheet = workbook["Research Sources"]
            self.assertEqual([sheet.cell(1, col).value for col in range(1, 15)], SPREADSHEET_COLUMNS)
            self.assertEqual(sheet.cell(2, 6).value, "2021–2023")
            audit = workbook["Data Period Audit"]
            self.assertEqual(audit.cell(2, 9).value, "Verified")
            self.assertEqual(audit.cell(2, 10).value, "Data were collected between 2021 and 2023.")
            qc_sheet = workbook["QC Audit & Coverage"]
            qc_values = [cell.value for row in qc_sheet.iter_rows() for cell in row]
            self.assertIn("India | Epidemiology | 2024", qc_values)
            workbook.close()
            document = Document(docx)
            text = "\n".join(p.text for p in document.paragraphs)
            tables = "\n".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
            self.assertIn("Data-period coverage by year", text)
            self.assertIn("2021–2023", tables)


if __name__ == "__main__":
    unittest.main()
