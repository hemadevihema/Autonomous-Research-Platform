"""
Stage 8: Organization Word Document Exporter (.docx)
=====================================================
Generates a formal, executive-grade research monograph in Microsoft Word (.docx) format.
Groups findings hierarchically by Country and then across all 9 Mandatory Research Sections.
Integrates evidence-backed insights, verified verbatim callouts, source tables, and QC disclosure.
"""
from pathlib import Path
from typing import List, Dict, Any
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

from config.settings import (
    DEFAULT_RESEARCH_SECTIONS, OUTPUT_DIR, MIN_PUBLICATION_YEAR,
    DATA_PERIOD_START_YEAR, DATA_PERIOD_END_YEAR,
)


class WordDocumentExporter:
    """Exporter producing structured .docx research reports adhering to organizational standards."""

    def export(
        self,
        project_id: str,
        topic: str,
        disease: str,
        countries: List[str],
        sections: List[str],
        validated_sources: List[Dict[str, Any]],
        insights: List[Dict[str, Any]],
        verbatims: List[Dict[str, Any]],
        qc_report: Dict[str, Any] = None,
        extraction_completeness: Dict[str, Dict[str, Any]] = None,
        min_publication_year: int = MIN_PUBLICATION_YEAR,
        data_period_start_year: int = DATA_PERIOD_START_YEAR,
        data_period_end_year: int = DATA_PERIOD_END_YEAR,
        review_sources: List[Dict[str, Any]] = None,
    ) -> str:
        """Construct and format the comprehensive Word document."""
        doc = Document()
        sources_by_id = {source.get("id"): source for source in validated_sources if source.get("id")}

        # Page margins
        sections_doc = doc.sections
        for s in sections_doc:
            s.top_margin = Inches(1.0)
            s.bottom_margin = Inches(1.0)
            s.left_margin = Inches(1.0)
            s.right_margin = Inches(1.0)

        # -------------------------------------------------------------
        # Document Title & Cover Header
        # -------------------------------------------------------------
        title_p = doc.add_paragraph()
        title_run = title_p.add_run(f"Strategic Healthcare Research Monograph:\n{disease.title()} Research Brief")
        title_run.font.name = "Calibri"
        title_run.font.size = Pt(24)
        title_run.font.bold = True
        title_run.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)  # Navy
        title_p.paragraph_format.space_after = Pt(12)

        meta_p = doc.add_paragraph()
        meta_run = meta_p.add_run(
            f"Project Reference: {project_id}  |  Target Geographies: {', '.join(countries)}  |  Publication Rule: >= {min_publication_year}\n"
            f"Preferred Data Collection Period: {data_period_start_year}-{data_period_end_year}\n"
            f"QC Completion Status: {qc_report.get('overall_status', 'PASSED')}  |  Verified Sources: {len(validated_sources)}"
        )
        meta_run.font.name = "Calibri"
        meta_run.font.size = Pt(10)
        meta_run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
        meta_p.paragraph_format.space_after = Pt(24)

        # -------------------------------------------------------------
        # Executive Summary & Methodology
        # -------------------------------------------------------------
        h1 = doc.add_heading("1. Executive Summary & Research Methodology", level=1)
        h1.style.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

        p_exec = doc.add_paragraph(
            f"This autonomous research report investigates the landscape of {disease} with specific focus on {', '.join(countries)}. "
            f"In accordance with rigorous research governance standards, candidate publications were audited against strict eligibility rules: "
            f"(1) Mandatory publication year {min_publication_year} or later, (2) Preferred data collection window "
            f"{data_period_start_year}-{data_period_end_year} with out-of-window periods flagged for review, and "
            f"(3) Independent string-level verification of verbatim quotations against primary source texts."
        )
        p_exec.style.font.name = "Calibri"
        p_exec.paragraph_format.space_after = Pt(14)

        # -------------------------------------------------------------
        # Findings Grouped by Country and the 9 Research Sections
        # -------------------------------------------------------------
        doc.add_heading("2. Evidence Synthesis by Country & Section", level=1)

        for country in countries:
            doc.add_heading(f"Country Profile: {country}", level=2)

            for sec_idx, section in enumerate(sections, start=1):
                doc.add_heading(f"{sec_idx}. {section}", level=3)

                # Filter sources for this country + section
                sec_sources = [
                    s for s in validated_sources
                    if s.get("country", "").lower() == country.lower() and section in s.get("sections_covered", [])
                ]

                # Filter insights and verbatims
                sec_insights = [
                    ins for ins in insights
                    if ins.get("country", "").lower() == country.lower() and ins.get("section") == section
                ]
                sec_verbatims = [
                    v for v in verbatims
                    if v.get("country", "").lower() == country.lower() and v.get("section") == section
                ]

                if not sec_sources:
                    p_empty = doc.add_paragraph(
                        f"Note: No primary sources satisfied the strict inclusion criteria for '{section}' during current scouting cycles. "
                        f"This area has been flagged as an active research gap for targeted follow-up."
                    )
                    p_empty.style.font.italic = True
                    p_empty.paragraph_format.space_after = Pt(10)
                    continue

                # Insights subsection
                p_ins_head = doc.add_paragraph()
                r_ins = p_ins_head.add_run("Key Analytical Insights:")
                r_ins.font.bold = True
                r_ins.font.size = Pt(11)

                for ins in sec_insights:
                    p_bullet = doc.add_paragraph(style="List Bullet")
                    p_bullet.add_run(f"{ins.get('insight_text', '')} ")
                    source = sources_by_id.get(ins.get("source_id"), {})
                    reference = source.get("source_title") or ins.get("source_id", "Source")
                    year = source.get("publication_year") or ins.get("publication_date")
                    if year:
                        reference += f" ({year})"
                    if source.get("source_link"):
                        reference += f" — {source['source_link']}"
                    ref_run = p_bullet.add_run(
                        f"[Source: {reference} | Location: {ins.get('location_reference', 'Source text')} | "
                        f"Data Period: {ins.get('data_period')} | Support: {ins.get('verification_status', 'unresolved')}]"
                    )
                    ref_run.font.size = Pt(9)
                    ref_run.font.italic = True

                # Verbatim callout blocks
                if sec_verbatims:
                    p_verb_head = doc.add_paragraph()
                    r_verb = p_verb_head.add_run("Verified Verbatim Quotations:")
                    r_verb.font.bold = True
                    r_verb.font.size = Pt(11)

                    for v in sec_verbatims:
                        quote_table = doc.add_table(rows=1, cols=1)
                        quote_table.alignment = WD_TABLE_ALIGNMENT.CENTER
                        cell = quote_table.cell(0, 0)
                        
                        shading = parse_xml(r'<w:shd {} w:fill="F0F4F8"/>'.format(nsdecls('w')))
                        cell._tc.get_or_add_tcPr().append(shading)
                        
                        p_callout = cell.paragraphs[0]
                        p_callout.add_run(f'“{v.get("exact_quote", "").replace(chr(34), "")}”\n').font.italic = True
                        
                        v_status = v.get("verification_status", "unresolved")
                        source = sources_by_id.get(v.get("source_id"), {})
                        source_ref = source.get("source_title") or v.get("source_id", "Source")
                        if source.get("source_link"):
                            source_ref += f" — {source['source_link']}"
                        p_cite = p_callout.add_run(
                            f"— Source: {source_ref} | Location: {v.get('location_reference')} | Status: {v_status}"
                        )
                        p_cite.font.size = Pt(9)
                        p_cite.font.bold = True
                        p_cite.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

                    doc.add_paragraph().paragraph_format.space_after = Pt(8)

        # -------------------------------------------------------------
        # Section 3: Evidence Limitations & Unresolved Gaps
        # -------------------------------------------------------------
        doc.add_heading("3. Evidence Limitations & Research Gaps", level=1)
        missing_cells = qc_report.get("missing_coverage", []) if qc_report else []
        if missing_cells:
            doc.add_paragraph(
                f"The following country-section domains currently lack direct primary literature meeting the configured publication cutoff ({min_publication_year}+):\n" +
                "\n".join([f" * {cell}" for cell in missing_cells])
            )
        else:
            doc.add_paragraph("All target country and research section combinations have verified primary evidence coverage.")

        doc.add_heading(
            f"Data-period coverage by year ({data_period_start_year}-{data_period_end_year})",
            level=2,
        )
        year_coverage = qc_report.get("year_wise_coverage", {}) if qc_report else {}
        if year_coverage:
            year_table = doc.add_table(rows=1, cols=6)
            year_table.alignment = WD_TABLE_ALIGNMENT.CENTER
            for index, header in enumerate(("Country", "Section", "Year", "Verified Sources", "Partial Sources", "Status")):
                year_table.rows[0].cells[index].text = header
                year_table.rows[0].cells[index].paragraphs[0].runs[0].font.bold = True
            for cell in year_coverage.values():
                row = year_table.add_row().cells
                row[0].text = str(cell.get("country", ""))
                row[1].text = str(cell.get("section", ""))
                row[2].text = str(cell.get("year", ""))
                row[3].text = str(cell.get("verified_sources_count", 0))
                row[4].text = str(cell.get("partial_sources_count", 0))
                row[5].text = str(cell.get("status", "No evidence"))
            gaps = qc_report.get("year_coverage_gaps", [])
            if gaps:
                doc.add_paragraph(
                    "Years without verified source coverage (including years with partial evidence only): "
                    + "; ".join(gaps)
                )
        else:
            doc.add_paragraph("Year-wise collection-period coverage was not available in the QC output.")

        review_sources = review_sources or []
        if review_sources:
            doc.add_heading("Sources retained for period or metadata review", level=2)
            review_table = doc.add_table(rows=1, cols=6)
            review_table.alignment = WD_TABLE_ALIGNMENT.CENTER
            for index, header in enumerate(("Title", "Publication Year", "Data Period", "Period Status", "Review Reason", "Source URL")):
                review_table.rows[0].cells[index].text = header
                review_table.rows[0].cells[index].paragraphs[0].runs[0].font.bold = True
            for source in review_sources:
                row = review_table.add_row().cells
                row[0].text = str(source.get("source_title", ""))
                row[1].text = str(source.get("publication_year") or "Unknown")
                row[2].text = str(source.get("data_period") or "Unknown")
                row[3].text = str(source.get("data_period_status") or "Unknown")
                row[4].text = "; ".join(source.get("validation_reasons", []) or [])
                row[5].text = str(source.get("source_link", ""))

        incomplete_sources = [
            (source_id, result)
            for source_id, result in (extraction_completeness or {}).items()
            if result.get("status") != "complete"
        ]
        if incomplete_sources:
            doc.add_paragraph(
                "Some sources did not meet the configured extraction completeness targets. "
                "Their available content was reported as retrieved; no extra findings or quotations were fabricated."
            )
            for source_id, result in incomplete_sources:
                doc.add_paragraph(
                    f"{source_id}: {result.get('insight_count', 0)} insight(s), "
                    f"{result.get('verbatim_count', 0)} verified quote(s); {result.get('reason', '')}",
                    style="List Bullet",
                )

        # -------------------------------------------------------------
        # Section 4: Comprehensive Reference Table
        # -------------------------------------------------------------
        doc.add_heading("4. Validated Source References", level=1)
        ref_table = doc.add_table(rows=1, cols=5)
        ref_table.alignment = WD_TABLE_ALIGNMENT.CENTER

        headers = ["Ref ID", "Title", "Publisher / Journal", "Year", "Data Period"]
        hdr_cells = ref_table.rows[0].cells
        for i, title in enumerate(headers):
            hdr_cells[i].text = title
            hdr_cells[i].paragraphs[0].runs[0].font.bold = True
            hdr_cells[i].paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            shading = parse_xml(r'<w:shd {} w:fill="1B365D"/>'.format(nsdecls('w')))
            hdr_cells[i]._tc.get_or_add_tcPr().append(shading)

        for s in validated_sources:
            row_cells = ref_table.add_row().cells
            row_cells[0].text = s.get("id", "")
            row_cells[1].text = s.get("source_title", "")[:60] + ("..." if len(s.get("source_title", "")) > 60 else "")
            row_cells[2].text = s.get("name_of_source", "")[:30]
            row_cells[3].text = str(s.get("publication_year", ""))
            row_cells[4].text = s.get("data_period", "")

        clean_name = f"Research_Report_{project_id}.docx"
        target_path = OUTPUT_DIR / clean_name
        doc.save(str(target_path))
        print(f"\n[Stage 8: Word Exporter] Word document generated successfully:")
        print(f"  [FILE SAVED] {target_path}")
        return str(target_path)
